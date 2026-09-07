"""Offline tests for the Excel renderer. No Azure, no network."""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "function_app"))

from render_report import build_workbook  # noqa: E402


def test_renderer_writes_original_prompt_blocks():
    summary = {
        "assessment_date": "2026-09-01",
        "applicant": {
            "Full Name(s)": "Alex Taylor",
            "Age(s)": "Not provided in binder",
            "Dependants": "Not provided in binder",
            "Address and living situation": "12 Example Street, Wellington 6011 (Renting)",
        },
        "accounts": [
            {
                "account_id": "deadbeef",
                "account_label": "01-0123-0456789-00",
                "institution": "ANZ",
                "account_type": "everyday",
                "period_start": "2026-07-01",
                "period_end": "2026-07-31",
                "days_covered": 31,
                "days_old": 32,
                "source_file": "ANZ.pdf",
                "is_statement": True,
                "conduct": {
                    "arrears_or_overdrawn": {"identified": False, "count": 0},
                    "late_dishonour_unarranged_fees": {"identified": False, "count": 0},
                    "limit_excess": {"identified": False, "count": 0},
                    "irregular_activity": {"identified": True, "count": 1},
                },
            }
        ],
        "document_index": [
            {
                "institution": "ANZ",
                "account_label": "01-0123-0456789-00",
                "account_type": "everyday",
                "period_start": "2026-07-01",
                "period_end": "2026-07-31",
                "days_covered": 31,
                "days_old": 32,
                "source_file": "ANZ.pdf",
                "conduct": {
                    "arrears_or_overdrawn": {"identified": False, "count": 0},
                    "late_dishonour_unarranged_fees": {"identified": False, "count": 0},
                    "limit_excess": {"identified": False, "count": 0},
                    "irregular_activity": {"identified": True, "count": 1},
                },
            }
        ],
        "income": [],
        "part1": [{"category": "Rent / board paid", "monthly_equivalent": 2400, "notes": "monthly"}],
        "part2": [],
        "part2_calculations": [
            {
                "merchant": "MERIDIAN",
                "category": "utilities",
                "dates_observed": ["2026-07-16"],
                "average_amount": 214.5,
                "assessed_frequency": "monthly",
                "calculated_monthly_impact": 214.5,
                "calculation_basis": "average amount",
            }
        ],
        "part3": [],
        "part4": [
            {
                "institution": "Westpac",
                "facility_type": "Credit card",
                "credit_limit": 6000,
                "current_balance": 388.4,
                "observed_repayment": 200,
                "frequency": "monthly",
                "evidence": "westpac.pdf",
            }
        ],
        "liability_evidence_status": "identified",
        "part5": {
            "top_merchants": [
                {
                    "merchant": "RENT LANDLORD",
                    "annual_spend": 28800,
                    "category": "rent_board_paid",
                    "calculation_basis": "average amount",
                }
            ],
            "high_frequency": [{"merchant": "PAK N SAVE", "hits": 4}],
            "underwriter_notes": [
                {"topic": "Account conduct", "note": "1 conduct exception", "requires_signoff": True}
            ],
        },
    }
    wb = load_workbook(BytesIO(build_workbook(summary)))
    names = set(wb.sheetnames)
    for required in (
        "1.1 Applicant",
        "1.2 Accounts",
        "1.3 Document index",
        "Part2 Calculations",
        "Part4 Liabilities",
        "Part5 Commentary",
        "_README",
    ):
        assert required in names, names

    applicant = {row[0].value: row[1].value for row in wb["1.1 Applicant"].iter_rows(min_row=2)}
    assert applicant["Full Name(s)"] == "Alex Taylor"
    assert "Renting" in applicant["Address and living situation"]

    account_cell = wb["1.2 Accounts"]["B2"].value
    assert account_cell == "01-0123-0456789-00"
    assert account_cell != "deadbeef"

    assert wb["1.3 Document index"]["H2"].value == "None identified"
    assert wb["1.3 Document index"]["K2"].value == "Identified (1)"
    assert wb["Part2 Calculations"]["A2"].value == "MERIDIAN"
    assert wb["Part4 Liabilities"]["A2"].value == "Westpac"
    assert wb["Part5 Commentary"]["B2"].value == 28800
    assert wb["Part5 Commentary"]["C2"].value == "rent_board_paid"
    part2_headers = [c.value for c in wb["Part2 Line items"][1]]
    assert "Direction" in part2_headers
    assert part2_headers[part2_headers.index("Amount") + 1] == "Direction"


def test_missing_is_business_warning_is_on_part5_not_only_json():
    summary = {
        "assessment_date": "2026-09-01",
        "part4": [],
        "liability_evidence_status": "no liability evidence identified",
        "audit": {"business_classification_missing": True},
        "part5": {
            "underwriter_notes": [
                {
                    "topic": "Business classification missing",
                    "note": "No classification in this run set is_business, so recommended living may include business spend. This was not assessed — not the same as finding no business spend.",
                    "requires_signoff": True,
                }
            ]
        },
    }
    wb = load_workbook(BytesIO(build_workbook(summary)))
    cells = [str(c.value or "") for row in wb["Part5 Commentary"].iter_rows() for c in row]
    blob = " ".join(cells)
    assert "UNDERWRITER WARNING" in blob
    assert "not assessed" in blob.lower()
    assert "is_business" in blob


def test_part2_writes_direction_column():
    summary = {
        "assessment_date": "2026-09-01",
        "part2": [
            {
                "date": "2026-07-02",
                "description": "Direct Credit MISS Y ZHANG",
                "amount": 40,
                "direction": "inflow",
                "frequency": "irregular",
                "include": "No",
                "category": "unclear",
                "source_file": "ANZ.pdf",
                "account": "ANZ",
                "exclusion_reason": "unclear",
                "needs_review": True,
            }
        ],
        "part4": [],
        "part5": {},
    }
    wb = load_workbook(BytesIO(build_workbook(summary)))
    headers = [c.value for c in next(wb["Part2 Line items"].iter_rows(min_row=1, max_row=1))]
    assert "Direction" in headers
    assert headers.index("Direction") == headers.index("Amount") + 1
    assert "Reason" in headers
    assert wb["Part2 Line items"]["D2"].value == "inflow"


def test_empty_part4_is_labelled_not_header_only():
    summary = {
        "assessment_date": "2026-09-01",
        "part4": [],
        "liability_evidence_status": "no liability evidence identified",
        "part5": {},
    }
    wb = load_workbook(BytesIO(build_workbook(summary)))
    assert "no liability evidence identified" in str(wb["Part4 Liabilities"]["A2"].value)


def test_evidence_gaps_render_as_their_own_block_and_vanish_when_empty():
    """The gaps block is written only when there are gaps to write.

    An empty block would read as "checked, nothing found" in exactly the same
    shape as "not checked at all", so absence has to mean absence.
    """

    gaps = [
        {
            "topic": "Account outside binder - ASB",
            "note": "3 transaction(s) name ASB, which has no statement in this binder.",
            "evidence": "2026-06-10 TFR TO ASB 300.00",
            "requires_signoff": True,
        },
        {
            "topic": "No power or gas in the file",
            "note": "No electricity or gas retailer appears in 61 statement-days.",
            "evidence": "no matching transactions",
            "requires_signoff": True,
        },
    ]
    summary = {"assessment_date": "2026-09-01", "part5": {"evidence_gaps": gaps}}
    ws = load_workbook(BytesIO(build_workbook(summary)))["Part5 Commentary"]
    rows = [[str(c.value or "") for c in row] for row in ws.iter_rows()]
    blob = " ".join(" ".join(r) for r in rows)

    header = next(r for r in rows if r and r[0] == "Evidence gaps")
    assert header[1:4] == ["What is missing", "Evidence", "Requires sign-off"], header
    assert "Account outside binder - ASB" in blob
    assert "No power or gas in the file" in blob
    assert "2026-06-10 TFR TO ASB 300.00" in blob
    signoff = [r[3] for r in rows if r and r[0].startswith(("Account outside", "No power"))]
    assert signoff == ["Yes", "Yes"], signoff

    # Same renderer, no gaps: the block must not appear at all.
    clean = {"assessment_date": "2026-09-01", "part5": {"evidence_gaps": []}}
    ws2 = load_workbook(BytesIO(build_workbook(clean)))["Part5 Commentary"]
    blob2 = " ".join(str(c.value or "") for row in ws2.iter_rows() for c in row)
    assert "Evidence gaps" not in blob2, "empty gaps must not write a bare header"
    assert "Underwriter audit notes" in blob2, "the rest of Part5 still renders"


def test_income_sheet_labels_turnover_gross_and_prints_the_assessable_total():
    summary = {
        "assessment_date": "2026-09-01",
        "income": [
            {"source": "ACME", "type": "salary_wages", "amount_observed": 4000,
             "frequency": "monthly", "monthly_equivalent": 4000, "evidence": "2026-06-03"},
            {"source": "Side business (unassessed) - 14 payers",
             "type": "side_business_gross_not_assessable", "amount_observed": 1400,
             "frequency": "irregular", "monthly_equivalent": 463.44,
             "gross_net": "GROSS RECEIPTS - not net profit. Excluded from assessable income; obtain business financials.",
             "evidence": "14 credits, 2026-06-01 to 2026-06-27"},
        ],
        "audit": {"assessable_income_monthly": 4000, "side_business_gross_monthly": 463.44},
        "part5": {},
    }
    ws = load_workbook(BytesIO(build_workbook(summary)))["1.4 Income"]
    rows = [[str(c.value or "") for c in row] for row in ws.iter_rows()]
    salary = next(r for r in rows if r[0] == "ACME")
    assert salary[5] == "unknown", salary
    side = next(r for r in rows if r[0].startswith("Side business"))
    assert "GROSS RECEIPTS" in side[5], side
    total = next(r for r in rows if r[0].startswith("ASSESSABLE INCOME"))
    assert total[4] == "4000", total
    assert "servicing" in total[5]


if __name__ == "__main__":
    test_renderer_writes_original_prompt_blocks()
    print("ok test_renderer_writes_original_prompt_blocks")
    test_empty_part4_is_labelled_not_header_only()
    print("ok test_empty_part4_is_labelled_not_header_only")
    test_missing_is_business_warning_is_on_part5_not_only_json()
    print("ok test_missing_is_business_warning_is_on_part5_not_only_json")
    test_part2_writes_direction_column()
    print("ok test_part2_writes_direction_column")
    test_evidence_gaps_render_as_their_own_block_and_vanish_when_empty()
    print("ok test_evidence_gaps_render_as_their_own_block_and_vanish_when_empty")
    test_income_sheet_labels_turnover_gross_and_prints_the_assessable_total()
    print("ok test_income_sheet_labels_turnover_gross_and_prints_the_assessable_total")
    print("ALL PASS")
