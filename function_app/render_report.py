# -*- coding: utf-8 -*-
"""Build lender-assessment.xlsx from compute_summary JSON. No Azure imports."""

from __future__ import annotations

import io
from typing import Any

from datetime import datetime

from openpyxl import Workbook

APPLICANT_FIELDS = (
    "Full Name(s)",
    "Age(s)",
    "Dependants",
    "Address and living situation",
)


def _business_cell(value: Any) -> str:
    flag = str(value or "no").strip().lower()
    if value is True or flag == "yes":
        return "Yes"
    if flag == "review":
        return "Review"
    return "No"


def _applicant_value(field: str, applicant: dict[str, Any], summary: dict[str, Any]) -> str:
    if field == "Assessment date":
        return applicant.get(field) or summary.get("assessment_date") or "Not provided in binder"
    for source in (applicant, summary.get("applicant") or {}):
        value = source.get(field)
        if value and value != "Not provided in binder":
            return value
    return "Not provided in binder"


def _as_excel_date(value: Any):
    if value in (None, ""):
        return None
    if hasattr(value, "year"):
        return value
    text = str(value)[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return value


def _conduct_cell(conduct: dict[str, Any] | None, key: str) -> str:
    if not isinstance(conduct, dict):
        return "NOT ASSESSABLE from extract"
    if conduct.get("status") == "Not returned by extractor":
        return "NOT ASSESSABLE from extract"
    block = conduct.get(key)
    if not isinstance(block, dict) or "identified" not in block:
        return "NOT ASSESSABLE from extract"
    if block.get("identified"):
        count = block.get("count")
        return f"Identified ({count})" if count not in (None, "") else "Identified"
    return "None identified"


def build_workbook(summary: dict[str, Any], applicant: dict[str, Any] | None = None) -> bytes:
    """Return xlsx bytes. Sheet content is owned by the engine, not the model."""
    applicant = applicant or {}
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("1.1 Applicant")
    ws.append(["Field", "Value"])
    for field in (*APPLICANT_FIELDS, "Assessment date"):
        ws.append([field, _applicant_value(field, applicant, summary)])

    ws = wb.create_sheet("1.2 Accounts")
    ws.append(["Institution", "Account", "Type", "Period", "Days covered", "Days old", "Source"])
    for acc in summary.get("accounts") or []:
        if acc.get("is_statement") is False:
            continue
        ws.append(
            [
                acc.get("institution"),
                acc.get("account_label") or "Not provided in binder",
                acc.get("account_type"),
                "%s to %s" % (acc.get("period_start"), acc.get("period_end")),
                acc.get("days_covered"),
                acc.get("days_old"),
                acc.get("source_file"),
            ]
        )

    ws = wb.create_sheet("1.3 Document index")
    ws.append(
        [
            "Institution",
            "Account",
            "Type",
            "Period",
            "Days covered",
            "Days old",
            "Source",
            "Arrears / overdrawn",
            "Late / dishonour / UOD fees",
            "Limit excess",
            "Irregular activity",
        ]
    )
    index = summary.get("document_index") or []
    if not index:
        index = [
            {
                **acc,
                "conduct": acc.get("conduct"),
            }
            for acc in (summary.get("accounts") or [])
            if acc.get("is_statement", True)
        ]
    for row in index:
        conduct = row.get("conduct") or {}
        ws.append(
            [
                row.get("institution"),
                row.get("account_label") or "Not provided in binder",
                row.get("account_type"),
                "%s to %s" % (row.get("period_start"), row.get("period_end")),
                row.get("days_covered"),
                row.get("days_old"),
                row.get("source_file"),
                _conduct_cell(conduct, "arrears_or_overdrawn"),
                _conduct_cell(conduct, "late_dishonour_unarranged_fees"),
                _conduct_cell(conduct, "limit_excess"),
                _conduct_cell(conduct, "irregular_activity"),
            ]
        )

    ws = wb.create_sheet("1.4 Income")
    ws.append(["Source", "Type", "Amount observed", "Frequency", "Monthly equivalent", "Gross/net", "Evidence"])
    for row in summary.get("income") or []:
        ws.append(
            [
                row.get("source"),
                row.get("type"),
                row.get("amount_observed"),
                row.get("frequency"),
                row.get("monthly_equivalent"),
                "unknown",
                row.get("evidence"),
            ]
        )

    ws = wb.create_sheet("Part1 Summary")
    ws.append(["Category", "Monthly equivalent", "Notes"])
    for row in summary.get("part1") or []:
        ws.append([row.get("category"), row.get("monthly_equivalent"), row.get("notes")])

    ws = wb.create_sheet("Part2 Line items")
    ws.append(
        [
            "Date",
            "Description",
            "Amount",
            "Direction",
            "Frequency",
            "Include",
            "Category",
            "Source file",
            "Account",
            "Exclusion reason",
            "Reason",
            "Needs review",
            "Is business",
        ]
    )
    for row in summary.get("part2") or []:
        ws.append(
            [
                _as_excel_date(row.get("date")),
                row.get("description"),
                row.get("amount"),
                row.get("direction"),
                row.get("frequency"),
                row.get("include"),
                row.get("category"),
                row.get("source_file"),
                row.get("account"),
                row.get("exclusion_reason"),
                row.get("reason"),
                "Yes" if row.get("needs_review") else "No",
                _business_cell(row.get("is_business")),
            ]
        )

    ws = wb.create_sheet("Part2 Calculations")
    ws.append(
        [
            "Merchant",
            "Category",
            "Dates observed",
            "Dominant amount",
            "Observed run-rate",
            "Assessment / frequency",
            "Calculated monthly impact",
            "Calculation basis",
        ]
    )
    for row in summary.get("part2_calculations") or []:
        dates = row.get("dates_observed") or []
        ws.append(
            [
                row.get("merchant"),
                row.get("category"),
                ", ".join(dates) if isinstance(dates, list) else dates,
                row.get("dominant_amount", row.get("average_amount")),
                row.get("observed_run_rate"),
                row.get("assessed_frequency"),
                row.get("calculated_monthly_impact"),
                row.get("calculation_basis"),
            ]
        )

    ws = wb.create_sheet("Part3 One-off")
    ws.append(["Date", "Description", "Amount", "Reason excluded"])
    for row in summary.get("part3") or []:
        ws.append([_as_excel_date(row.get("date")), row.get("description"), row.get("amount"), row.get("reason")])

    ws = wb.create_sheet("Part4 Liabilities")
    ws.append(
        [
            "Institution / lender",
            "Facility type",
            "Credit limit",
            "Current balance",
            "Observed repayment",
            "Frequency",
            "Evidence",
        ]
    )
    part4 = summary.get("part4")
    if part4 is None:
        ws.append(["OUTPUT_CONTRACT_FAILED", "part4 absent", "", "", "", "", ""])
    elif not part4:
        status = summary.get("liability_evidence_status") or "no liability evidence identified"
        ws.append([status, "", "", "", "", "", ""])
    else:
        for row in part4:
            ws.append(
                [
                    row.get("institution"),
                    row.get("facility_type"),
                    row.get("credit_limit"),
                    row.get("current_balance"),
                    row.get("observed_repayment"),
                    row.get("frequency"),
                    row.get("evidence"),
                ]
            )

    ws = wb.create_sheet("Part5 Commentary")
    missing_business = bool((summary.get("audit") or {}).get("business_classification_missing"))
    if not missing_business:
        missing_business = any(
            (row.get("topic") or "") == "Business classification missing"
            for row in (summary.get("part5") or {}).get("underwriter_notes") or []
        )
    if missing_business:
        ws.append(
            [
                "UNDERWRITER WARNING",
                "No classification in this run set is_business, so recommended living may include business spend. This was not assessed — not the same as finding no business spend.",
            ]
        )
        ws.append([])
    ws.append(["Merchant", "Annual spend", "Category", "Calculation basis"])
    for row in (summary.get("part5") or {}).get("top_merchants") or []:
        ws.append(
            [
                row.get("merchant"),
                row.get("annual_spend"),
                row.get("category"),
                row.get("calculation_basis"),
            ]
        )
    ws.append([])
    ws.append(["High-frequency merchants (>=3)", "Hits"])
    for row in (summary.get("part5") or {}).get("high_frequency") or []:
        ws.append([row.get("merchant"), row.get("hits")])
    ws.append([])
    gaps = (summary.get("part5") or {}).get("evidence_gaps") or []
    if gaps:
        ws.append(["Evidence gaps", "What is missing", "Evidence", "Requires sign-off"])
        for row in gaps:
            ws.append(
                [
                    row.get("topic"),
                    row.get("note"),
                    row.get("evidence"),
                    "Yes" if row.get("requires_signoff") else "No",
                ]
            )
        ws.append([])
    ws.append(["Underwriter audit notes", "Note", "Requires sign-off"])
    notes = (summary.get("part5") or {}).get("underwriter_notes") or []
    if not notes:
        ws.append(["File quality", "Not returned by tool", "Yes"])
    else:
        for row in notes:
            ws.append(
                [
                    row.get("topic"),
                    row.get("note"),
                    "Yes" if row.get("requires_signoff") else "No",
                ]
            )

    ws = wb.create_sheet("_README")
    ws.append(["Block", "Source"])
    ws.append(["Applicant / accounts / document index / conduct", "extract_and_normalize + compute_summary"])
    ws.append(["Part 1-5 figures", "compute_summary"])
    ws.append(["Classification", "model; no arithmetic"])
    ws.append(["MODEL DRAFT", "A human underwriter signs this off"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
